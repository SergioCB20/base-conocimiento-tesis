"""Arma la v2 de la tesis (.docx) a partir de capitulos/*.md.

Uso:  pip install pypandoc_binary python-docx pymupdf
      python build_docx.py
Los diagramas (diagramas/figNN.png) se generaron desde los bloques Mermaid del Cap. 4 con
@mermaid-js/mermaid-cli; si cambian los diagramas, regenerarlos antes (figNN.mmd están aquí).
Archivos intermedios (tesis_v2.md, reference.docx) se escriben en esta carpeta.
"""
import io, re, os, shutil, subprocess, pypandoc
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:/Users/a2020/Documents/Tesis/Base de conocimiento"
CAP = BASE + "/capitulos"
SCR = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = BASE + "/informes/Entrega 2"
OUT = OUT_DIR + "/20202117_SergioChumbimuni_LuisVives_E2_v2.docx"
LOGO = SCR + "/logo_pucp.png"
MMD = SCR + "/diagramas"
BR = '`<w:r><w:br/></w:r>`{=openxml}'
PAGEBREAK = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'

def read(n):
    return io.open(f"{CAP}/{n}", encoding="utf-8").read()

def clean(t):
    # comentarios PENDIENTE -> texto visible en rojo
    t = re.sub(r"<!--\s*PENDIENTE:(.*?)-->",
               lambda m: "[PENDIENTE:" + m.group(1).strip().replace("[", "(").replace("]", ")") + ']{custom-style="Pendiente"}',
               t, flags=re.S)
    t = re.sub(r"<!--.*?-->\n?", "", t, flags=re.S)                     # otros comentarios
    t = re.sub(r"(?m)^> \*\*Nota de transcripci[oó]n.*?(?:\n>.*)*\n?", "", t)  # notas de transcripción
    t = t.replace("<br>", BR)
    t = t.replace("](../imagenes/", "](" + BASE + "/imagenes/")
    return t

# ---------- portada ----------
cover = f"""
::: {{custom-style="Portada"}}
![]({LOGO}){{width=3.5cm}}

**PONTIFICIA UNIVERSIDAD CATÓLICA DEL PERÚ**

**FACULTAD DE CIENCIAS E INGENIERÍA**

**ESPECIALIDAD DE INGENIERÍA INFORMÁTICA**



**Desarrollo de una Arquitectura MLOps para la Validación Continua de Equidad (Fairness) como Prueba de Calidad (QA) en Modelos de Scoring Crediticio**



Tesis para obtener el título profesional de Ingeniero Informático



**AUTOR:** Sergio Alonso Chumbimuni Bustamante

**ASESOR:** Luis Vives Garnique



Lima, Setiembre, 2026

Versión 2 (borrador de revisión)
:::
"""
toc = """
::: {custom-style="TOC Heading"}
Índice
:::

```{=openxml}
<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>Índice: haga clic derecho aquí y elija «Actualizar campos» (o presione F9).</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
```
"""

parts = [cover, PAGEBREAK, toc]
for n in ["01-generalidades.md", "02-marco-conceptual-teorico-legal.md", "03-estado-del-arte.md",
          "04-diseno.md", "05-implementacion.md"]:
    t = clean(read(n))
    if n.startswith("03"):
        t = t.replace("**Figura 2: Arquitectura del sistema (Nivel 2 de Diagrama C4)**",
                      '**Figura 2: Arquitectura del sistema (Nivel 2 de Diagrama C4)**\n\n[PENDIENTE: esta figura corresponde al diseño anterior; reemplazarla por la Figura 5 o retirarla.]{custom-style="Pendiente"}')
    if n.startswith("04"):
        i = [4]
        def repl(m):
            k = i[0]; i[0] += 1
            import pymupdf
            px = pymupdf.Pixmap(f"{MMD}/fig{k:02d}.png")
            attr = "height=20cm" if px.height / px.width > 20 / 15.5 else "width=15.5cm"
            return f"![]({MMD}/fig{k:02d}.png){{{attr}}}\n"
        t = re.sub(r"```mermaid\n.*?```\n", repl, t, flags=re.S)
        t = t.replace("**Figura 4: Diagrama de Contexto (C4, Nivel 1)**",
                      '**Figura 4: Diagrama de Contexto (C4, Nivel 1)**\n\n[PENDIENTE: las Figuras 4 a 9 son diagramas provisionales generados automáticamente; redibujarlos en Lucidchart.]{custom-style="Pendiente"}')
    parts.append(t)

# referencias: viñetas -> párrafos con sangría francesa
ref = clean(read("06-referencias.md"))
lines = []
for l in ref.splitlines():
    if l.startswith("- "):
        lines.append('::: {custom-style="Bibliografia"}\n' + l[2:] + "\n:::\n")
    else:
        lines.append(l)
parts.append("\n".join(lines))
parts.append(clean(read("07-anexos.md")))
md = "\n\n".join(parts)
io.open(SCR + "/tesis_v2.md", "w", encoding="utf-8").write(md)

# ---------- documento de referencia (estilos) ----------
ref_path = SCR + "/reference.docx"
subprocess.run([pypandoc.get_pandoc_path(), "-o", ref_path, "--print-default-data-file", "reference.docx"], check=True)
rd = Document(ref_path)
sec = rd.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.54))
st = rd.styles

def font(s, size=None, bold=None, color=RGBColor(0, 0, 0), name="Times New Roman", italic=None):
    s.font.name = name
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    if size: s.font.size = Pt(size)
    if bold is not None: s.font.bold = bold
    if italic is not None: s.font.italic = italic
    if color is not None: s.font.color.rgb = color

for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
    if name in [x.name for x in st]:
        s = st[name]; font(s, 14)
        pf = s.paragraph_format; pf.line_spacing = 1.15; pf.space_after = Pt(6)
        if name in ("Body Text", "First Paragraph"):
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size in (("Heading 1", 18), ("Heading 2", 16), ("Heading 3", 14), ("Heading 4", 14)):
    s = st[name]; font(s, size, bold=True, italic=False)
    s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(6)
st["Heading 1"].paragraph_format.page_break_before = True
for name in ("Title", "Subtitle"):
    font(st[name], 18, bold=True)
if "Source Code" in [x.name for x in st]:
    font(st["Source Code"], 10, name="Consolas")
if "Verbatim Char" in [x.name for x in st]:
    font(st["Verbatim Char"], 10, name="Consolas")
if "Hyperlink" in [x.name for x in st]:
    st["Hyperlink"].font.color.rgb = RGBColor(0x1F, 0x4E, 0x9A)

def pstyle(name, align=None, size=14, bold=None, hanging=None):
    s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH) if name not in [x.name for x in st] else st[name]
    s.base_style = st["Normal"]; font(s, size, bold=bold)
    if align is not None: s.paragraph_format.alignment = align
    if hanging:
        s.paragraph_format.left_indent = Cm(hanging); s.paragraph_format.first_line_indent = Cm(-hanging)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return s
pstyle("Portada", WD_ALIGN_PARAGRAPH.CENTER, 14)
pstyle("Bibliografia", size=12, hanging=1.27)
s = pstyle("TOC Heading", WD_ALIGN_PARAGRAPH.LEFT, 18, bold=True)
cs = st.add_style("Pendiente", WD_STYLE_TYPE.CHARACTER) if "Pendiente" not in [x.name for x in st] else st["Pendiente"]
cs.font.color.rgb = RGBColor(0xC0, 0x00, 0x00); cs.font.bold = True; cs.font.size = Pt(12)
rd.save(ref_path)

# ---------- conversión ----------
os.makedirs(OUT_DIR, exist_ok=True)
pypandoc.convert_file(SCR + "/tesis_v2.md", "docx", format="markdown-implicit_figures+raw_attribute",
                      outputfile=OUT, extra_args=["--reference-doc=" + ref_path, "--resource-path=" + BASE])

# ---------- post-proceso ----------
TBLPR_ORDER = ["tblStyle","tblpPr","tblOverlap","bidiVisual","tblStyleRowBandSize","tblStyleColBandSize","tblW","jc","tblCellSpacing","tblInd","tblBorders","shd","tblLayout","tblCellMar","tblLook","tblCaption","tblDescription"]
def put_ordered(parent, el, order):
    name = el.tag.split("}")[1]
    old = parent.find(qn("w:" + name))
    if old is not None: parent.remove(old)
    idx = order.index(name)
    for i, ch in enumerate(list(parent)):
        cn = ch.tag.split("}")[1]
        if cn in order and order.index(cn) > idx:
            ch.addprevious(el); return
    parent.append(el)

doc = Document(OUT)
def set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000")
        b.append(e)
    put_ordered(tblPr, b, TBLPR_ORDER)
from docx.shared import Emu
TOTAL_CM = 15.9
def fit_widths(tbl):
    ncols = len(tbl.columns)
    lens = [0.0]*ncols
    for row in tbl.rows:
        for j, cell in enumerate(row.cells[:ncols]):
            words = cell.text.split()
            longest = max((len(w) for w in words), default=1)
            lens[j] = max(lens[j], min(len(cell.text), 400) ** 0.6, longest * 0.9)
    tot = sum(lens) or 1
    widths = [max(TOTAL_CM * l / tot, 1.4) for l in lens]
    k = TOTAL_CM / sum(widths); widths = [w*k for w in widths]
    tblPr = tbl._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); put_ordered(tblPr, lay, TBLPR_ORDER)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None: tw = OxmlElement("w:tblW"); put_ordered(tblPr, tw, TBLPR_ORDER)
    tw.set(qn("w:w"), str(int(TOTAL_CM/2.54*1440))); tw.set(qn("w:type"), "dxa")
    grid = tbl._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[j]/2.54*1440)))
    for row in tbl.rows:
        for j, cell in enumerate(row.cells[:ncols]):
            cell.width = Cm(widths[j])
    return ncols
for tbl in doc.tables:
    set_borders(tbl)
    ncols = fit_widths(tbl)
    fsize = 9.5 if ncols >= 6 else 11
    for i, row in enumerate(tbl.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(fsize)
                    if i == 0: r.font.bold = True
    # repetir encabezado
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
cap = re.compile(r"^(Figura|Tabla)\s+\d+")
for p in doc.paragraphs:
    txt = p.text.strip()
    if cap.match(txt) and len(txt) < 200:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = False
    if any(r._element.find(qn("w:drawing")) is not None for r in p.runs) and p.style.name != "Portada":
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
# que Word actualice el índice al abrir
settings = doc.settings.element
u = OxmlElement("w:updateFields"); u.set(qn("w:val"), "true")
after = ["hdrShapeDefaults","footnotePr","endnotePr","compat","docVars","rsids","mathPr","attachedSchema","themeFontLang","clrSchemeMapping","doNotIncludeSubdocsInStats","doNotAutoCompressPictures","forceUpgrade","captions","readModeInkLockDown","smartTagType","schemaLibrary","shapeDefaults","doNotEmbedSmartTags","decimalSymbol","listSeparator"]
for ch in list(settings):
    if ch.tag.split("}")[1] in after:
        ch.addprevious(u); break
else:
    settings.append(u)
for sect in doc.element.body.iter(qn("w:sectPr")):
    pm = sect.find(qn("w:pgMar"))
    if pm is not None:
        for a, v in (("w:header", "708"), ("w:footer", "708"), ("w:gutter", "0")):
            if pm.get(qn(a)) is None: pm.set(qn(a), v)
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
for rpr in doc.element.body.iter(M + "rPr"):
    if rpr.find(M + "nor") is not None:
        for s in rpr.findall(M + "sty"):
            rpr.remove(s)
doc.core_properties.title = "Desarrollo de una Arquitectura MLOps para la Validación Continua de Equidad (Fairness) - v2"
doc.core_properties.author = "Sergio Alonso Chumbimuni Bustamante"
doc.save(OUT)
print("OK", OUT, os.path.getsize(OUT))
